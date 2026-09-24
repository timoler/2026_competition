"""Independent arithmetic verification of Q1 (recomputes geometry from raw data).

Does NOT treat main.py's exported legs / payload limits as ground truth.  It
independently reconstructs UTM coordinates, leg distances, DEM max elevation and
cruise altitude from the raw node/DEM data via the shared geometry module, then
checks main.py's exports match, and independently re-derives max payload,
per-trip mass/volume, energy, time, reserve and box coverage.
"""
import csv
import json
import math
import sys
from itertools import permutations
from pathlib import Path

import openpyxl
from pyproj import Transformer

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "Q2" / "code"))
from prepare_data import Terrain  # reuse Q2's verified UTM-straight-line DEM traversal

P = Path(__file__).resolve().parents[1]
R = P / "results"
S = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / "data/raw/D题"
BASE = S / "数据/无人机应急物资运输基础数据"


def read(n):
    return list(csv.DictReader((R / n).open(encoding="utf-8-sig")))


def source(name):
    found = list(S.rglob(name))
    if len(found) != 1:
        raise ValueError(f"Expected exactly one {name}: {found}")
    return found[0]


checks = []
def check(name, ok, detail=""):
    checks.append((name, bool(ok), detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {name} {detail}")


# --- 1. raw nodes + independent UTM coordinates -----------------------------
w = openpyxl.load_workbook(BASE / "调度中心与服务区.xlsx", data_only=True)
nr = list(w.active.values)
nodes = {r[0]: dict(lon=r[2], lat=r[3], z=r[4]) for r in [nr[2]] + nr[6:] if r[0]}
assert len(nodes) == 16
utm = Transformer.from_crs(4326, 32649, always_xy=True)
for k, n in nodes.items():
    n["x_m"], n["y_m"] = utm.transform(n["lon"], n["lat"])
    n["work_z"] = n["z"] + (0 if k == "O01" else 30)

# exported nodes comparison
exported_nodes = {x["id"]: x for x in read("q1_nodes.csv")}
utm_ok = all(abs(float(exported_nodes[k]["x_m"]) - nodes[k]["x_m"]) < 1e-6 and
             abs(float(exported_nodes[k]["y_m"]) - nodes[k]["y_m"]) < 1e-6 for k in nodes)
check("utm_coordinates", utm_ok)

# --- 2. independent leg geometry (shared Terrain, same as Q2) ---------------
terrain = Terrain(source("镇龙乡及周边30米DEM.mat"))
indep = {}
for i, j in permutations(nodes, 2):
    ni, nj = nodes[i], nodes[j]
    peak, count = terrain.maximum((ni["x_m"], ni["y_m"]), (nj["x_m"], nj["y_m"]))
    H = peak + 50
    indep[(i, j)] = dict(distance_m=math.hypot(nj["x_m"] - ni["x_m"], nj["y_m"] - ni["y_m"]),
                         terrain_max_m=peak, cruise_m=H,
                         up_m=H - ni["work_z"], down_m=H - nj["work_z"], cells=count)

exported_legs = {(x["start"], x["end"]): x for x in read("q1_legs.csv")}
max_dd = max_dem = 0.0
leg_ok = True
for k, ind in indep.items():
    ex = exported_legs[k]
    dd = abs(float(ex["distance_m"]) - ind["distance_m"])
    dem = abs(float(ex["terrain_max_m"]) - ind["terrain_max_m"])
    max_dd = max(max_dd, dd); max_dem = max(max_dem, dem)
    if dd > 1e-6 or dem > 1e-6 or abs(float(ex["cruise_m"]) - ind["cruise_m"]) > 1e-6 \
       or abs(float(ex["up_m"]) - ind["up_m"]) > 1e-6 or abs(float(ex["down_m"]) - ind["down_m"]) > 1e-6:
        leg_ok = False
check("leg_geometry_independent", leg_ok, f"max dist diff={max_dd:.2e} m, max dem diff={max_dem:.2e} m")

# --- 3. types / boxes -------------------------------------------------------
w = openpyxl.load_workbook(BASE / "运输无人机数据.xlsx", data_only=True)
keys = ["id", "name", "mass", "Q", "V", "speed", "L0", "LF", "battery", "reserve_pct",
        "prep", "load", "handover", "perbox", "up", "down", "eta", "downeta"]
models = {r[0]: dict(zip(keys, r)) for r in list(w.active.values)[2:5]}
w = openpyxl.load_workbook(BASE / "物资需求与配送时限.xlsx", data_only=True)
boxes = {r[0]: r for r in list(w["逐箱货箱清单"].values)[1:]}
assert len(boxes) == 80


def leg_perf(i, j, k, q):
    g = models[k]; l = indep[(i, j)]
    L = g["L0"] - (g["L0"] - g["LF"]) * (q / g["Q"]) ** 1.5
    eh = g["battery"] * l["distance_m"] / L
    eu = (g["mass"] + q) * 9.80665 * l["up_m"] / (g["eta"] * 3.6e6)
    t = l["up_m"] / g["up"] + l["distance_m"] / g["speed"] + l["down_m"] / g["down"]
    return t, eh + eu


def trip_energy_time(site, k, q, count):
    g = models[k]
    t1, e1 = leg_perf("O01", site, k, q)
    t2, e2 = leg_perf(site, "O01", k, 0)
    prep = g["prep"] + g["load"] * count
    handover = g["handover"] + g["perbox"] * count
    return t1 + t2 + prep + handover, e1 + e2


def capacity(site, k, r):
    g = models[k]; limit = (1 - r) * g["battery"]
    if trip_energy_time(site, k, 0, 0)[1] > limit:
        return None
    if trip_energy_time(site, k, float(g["Q"]), 0)[1] <= limit:
        return float(g["Q"])
    lo, hi = 0.0, float(g["Q"])
    for _ in range(80):
        mid = (lo + hi) / 2
        if trip_energy_time(site, k, mid, 0)[1] <= limit:
            lo = mid
        else:
            hi = mid
    return lo


# --- 4. max payload comparison ---------------------------------------------
exported_cap = {(x["reserve"], x["site"], x["model"]): x for x in read("q1_max_payload_sensitivity.csv")}
cap_ok = True
for (r, site, k), ex in exported_cap.items():
    ind = capacity(site, k, float(r))
    ex_reach = ex["reachable"] == "True"
    ind_reach = ind is not None
    if ex_reach != ind_reach:
        cap_ok = False
    elif ind is not None and abs(float(ex["max_payload_kg"]) - ind) > 1e-6:
        cap_ok = False
check("max_payload_independent", cap_ok)

# --- 5. batches: box coverage, mass/volume, energy, time, reserve -----------
plan_files = [R / "q1_batches_default.csv"] + sorted(R.glob("q1_batches_reserve_*.csv")) + sorted(R.glob("q1_batches_priority_*.csv"))
checked = 0
all_ok = True
for f in plan_files:
    reserve = float(f.name.rsplit("_", 1)[1].replace(".csv", "")) / 100 if "reserve_" in f.name else 0.2
    assigned = []
    for x in read(f.name):
        m = models[x["model"]]
        cargo = [boxes[i] for i in x["box_ids"].split(";")]
        assigned.extend(r_[0] for r_ in cargo)
        mass = sum(z[3] for z in cargo); vol = sum(z[4] for z in cargo)
        if not (all(z[1] == x["site"] for z in cargo) and mass <= m["Q"] and vol <= m["V"] + 1e-12):
            all_ok = False
        op, energy = trip_energy_time(x["site"], x["model"], mass, len(cargo))
        flight = leg_perf("O01", x["site"], x["model"], mass)[0] + leg_perf(x["site"], "O01", x["model"], 0)[0]
        prep = m["prep"] + m["load"] * len(cargo)
        handover = m["handover"] + m["perbox"] * len(cargo)
        if any(not math.isclose(float(x[field]), expected, abs_tol=1e-7)
               for field, expected in (("flight_s", flight), ("preparation_s", prep), ("handover_s", handover))):
            all_ok = False
        if not math.isclose(energy, float(x["energy_kwh"]), abs_tol=1e-9):
            all_ok = False
        if not math.isclose(op, float(x["operation_s"]), abs_tol=1e-7):
            all_ok = False
        if not math.isclose(100 * (1 - energy / m["battery"]), float(x["soc_pct"]), abs_tol=1e-7):
            all_ok = False
        if not math.isclose(mass, float(x["mass_kg"]), abs_tol=1e-9) or not math.isclose(vol, float(x["volume_m3"]), abs_tol=1e-12):
            all_ok = False
        if not energy <= (1 - reserve) * m["battery"] + 1e-9:
            all_ok = False
        checked += 1
    if not (len(assigned) == 80 and set(assigned) == set(boxes)):
        all_ok = False
check("batches_mass_volume_energy_time_reserve", all_ok, f"{checked} trip records")

# --- 6. submission rows -----------------------------------------------------
sub = read("q1_submission_rows.csv")
base = read("q1_batches_default.csv")
sub_ok = (len(sub) == len(base) and
          all(s["架次编号"] == b["trip_id"] and s["服务区编号"] == b["site"] and
              s["机型编号"] == b["model"] and s["货箱编号列表"] == b["box_ids"]
              for s, b in zip(sub, base)))
check("submission_rows", sub_ok)
field_map = {"总质量（kg）": "mass_kg", "总体积（m³）": "volume_m3",
             "往返时间（s）": "operation_s", "架次能耗（kWh）": "energy_kwh", "返航SOC（%）": "soc_pct"}
check("submission_numeric_fields", len(sub) == len(base) and all(
    math.isclose(float(s[shown]), float(b[source_field]), abs_tol=1e-8)
    for s, b in zip(sub, base) for shown, source_field in field_map.items()))
summary = [r for r in read("q1_model_comparison.csv")
           if float(r["reserve"]) == 0.2 and r["priority"] == "时间_架次_能耗"]
check("default_summary", len(summary) == 1 and int(summary[0]["trips"]) == len(base) and all(
    math.isclose(float(summary[0][field]), math.fsum(float(b[field]) for b in base), abs_tol=1e-8)
    for field in ("energy_kwh", "operation_s", "flight_s")))

result = dict(status="PASS" if all(ok for _, ok, _ in checks) else "FAIL",
              checks=[{"check": n, "pass": ok, "detail": d} for n, ok, d in checks])
(R / "q1_verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print("\nOVERALL:", result["status"])
sys.exit(0 if result["status"] == "PASS" else 1)
